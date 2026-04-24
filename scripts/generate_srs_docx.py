"""Convert SRS.md to SRS.docx with proper Word formatting."""
import re
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def set_cell_bg(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def parse_table(lines):
    rows = []
    for line in lines:
        if re.match(r"^\s*\|[-:| ]+\|\s*$", line):
            continue
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        rows.append(cells)
    return rows


def add_table_to_doc(doc, rows):
    if not rows:
        return
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j, cell_text in enumerate(row):
            cell = table.cell(i, j)
            cell.text = cell_text
            para = cell.paragraphs[0]
            para.runs[0].font.size = Pt(9) if i > 0 else Pt(9)
            if i == 0:
                para.runs[0].bold = True
                set_cell_bg(cell, "2E4057")
                para.runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            elif i % 2 == 0:
                set_cell_bg(cell, "F2F2F2")


def add_heading(doc, text, level):
    heading = doc.add_heading(text, level=level)
    run = heading.runs[0] if heading.runs else heading.add_run(text)
    if level == 1:
        run.font.color.rgb = RGBColor(0x1A, 0x27, 0x44)
    elif level == 2:
        run.font.color.rgb = RGBColor(0x2E, 0x40, 0x57)
    elif level == 3:
        run.font.color.rgb = RGBColor(0x3D, 0x59, 0x7A)


def process_inline(run, text):
    """Handle **bold** text in a run."""
    run.text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)


def add_paragraph_with_inline(doc, text, style=None):
    if style:
        para = doc.add_paragraph(style=style)
    else:
        para = doc.add_paragraph()
    # Split on **bold**
    parts = re.split(r"(\*\*[^*]+\*\*)", text)
    for part in parts:
        if part.startswith("**") and part.endswith("**"):
            run = para.add_run(part[2:-2])
            run.bold = True
        else:
            para.add_run(part)
    return para


def main():
    with open("/Users/mebankhyriem/Desktop/asset-it/SRS.md", encoding="utf-8") as f:
        content = f.read()

    doc = Document()

    # Page margins
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1.2)
        section.right_margin = Inches(1.2)

    # Default font
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10)

    lines = content.split("\n")
    i = 0
    in_code = False
    code_lines = []
    in_table = False
    table_lines = []

    while i < len(lines):
        line = lines[i]

        # Code block
        if line.strip().startswith("```"):
            if not in_code:
                in_code = True
                code_lines = []
                i += 1
                continue
            else:
                in_code = False
                code_text = "\n".join(code_lines)
                para = doc.add_paragraph()
                run = para.add_run(code_text)
                run.font.name = "Courier New"
                run.font.size = Pt(8)
                para.paragraph_format.left_indent = Inches(0.3)
                # Light grey background via direct XML shading on paragraph
                pPr = para._p.get_or_add_pPr()
                shd = OxmlElement("w:shd")
                shd.set(qn("w:val"), "clear")
                shd.set(qn("w:color"), "auto")
                shd.set(qn("w:fill"), "F4F4F4")
                pPr.append(shd)
                code_lines = []
                i += 1
                continue

        if in_code:
            code_lines.append(line)
            i += 1
            continue

        # Table row detection
        if line.strip().startswith("|"):
            if not in_table:
                in_table = True
                table_lines = []
            table_lines.append(line)
            i += 1
            # Check if next line is also a table row
            if i < len(lines) and lines[i].strip().startswith("|"):
                continue
            else:
                in_table = False
                rows = parse_table(table_lines)
                add_table_to_doc(doc, rows)
                doc.add_paragraph()
                table_lines = []
                continue

        if in_table and not line.strip().startswith("|"):
            in_table = False
            rows = parse_table(table_lines)
            add_table_to_doc(doc, rows)
            doc.add_paragraph()
            table_lines = []

        stripped = line.strip()

        # Headings
        h_match = re.match(r"^(#{1,4})\s+(.*)", stripped)
        if h_match:
            level = len(h_match.group(1))
            text = h_match.group(2)
            # Remove markdown bold from heading text
            text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
            add_heading(doc, text, min(level, 4))
            i += 1
            continue

        # Horizontal rule
        if re.match(r"^---+$", stripped):
            doc.add_paragraph()
            i += 1
            continue

        # Bullet list
        if stripped.startswith("- "):
            text = stripped[2:]
            para = doc.add_paragraph(style="List Bullet")
            parts = re.split(r"(\*\*[^*]+\*\*)", text)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    run = para.add_run(part[2:-2])
                    run.bold = True
                else:
                    para.add_run(part)
            i += 1
            continue

        # Empty line
        if not stripped:
            i += 1
            continue

        # Normal paragraph
        add_paragraph_with_inline(doc, stripped)
        i += 1

    doc.save("/Users/mebankhyriem/Desktop/asset-it/SRS.docx")
    print("SRS.docx generated successfully.")


if __name__ == "__main__":
    main()
